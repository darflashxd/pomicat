"""Build fictional English Word Online fixtures. No network or Office account access.

Requires python-docx and Pillow (test tooling only). Run:
  python tools/build_word_corpus.py --output <directory>
  python tools/build_word_corpus.py --self-check
"""
import argparse
import hashlib
import json
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor
from PIL import Image, ImageDraw, ImageFont

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
NAME = "Avery Example"
EMAIL = "avery@example.invalid"
MARKER = "Synthetic demonstration document"


def font(size):
    try:
        return ImageFont.truetype("DejaVuSans.ttf", size)
    except OSError:
        return ImageFont.load_default()


def screenshot(path):
    image = Image.new("RGB", (1200, 420), "#f8fafc")
    draw = ImageDraw.Draw(image)
    lines = [MARKER, "Support case: DEMO-ONLY-0042", f"Contact: {EMAIL}",
             "Password: DEMO_ONLY_NOT_A_REAL_SECRET", "Public status: Investigation complete"]
    for index, text in enumerate(lines):
        draw.text((32, 22 + index * 74), text, fill="#17202a", font=font(28))
    image.save(path)
    # Coordinates are authored coarse ground truth, not OCR detections.
    return [{"kind": "email", "box": [24, 164, 1176, 226]},
            {"kind": "credential", "box": [24, 238, 1176, 300]}]


def equation(doc):
    paragraph = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    fraction = OxmlElement("m:f")
    for tag, text in [("m:num", "a+b"), ("m:den", "c")]:
        element = OxmlElement(tag)
        run = OxmlElement("m:r")
        value = OxmlElement("m:t")
        value.text = text
        run.append(value)
        element.append(run)
        fraction.append(element)
    math.append(fraction)
    paragraph._p.append(math)


def hidden_revision(doc):
    paragraph = doc.add_paragraph()
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "1")
    deleted.set(qn("w:author"), "Synthetic Author")
    run = OxmlElement("w:r")
    text = OxmlElement("w:delText")
    text.text = "OLD_DELETED_PRIVATE_SENTINEL"
    run.append(text)
    deleted.append(run)
    paragraph._p.append(deleted)


def build(output):
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise FileExistsError(f"Refusing to write into non-empty directory: {output}")
    image_path = output / "synthetic-support.png"
    regions = screenshot(image_path)
    specs = [
        ("01-onboarding", "Employee onboarding letter", "Please welcome the employee listed below.", "text"),
        ("02-consulting", "International consulting agreement", "This fictional agreement uses English prose and test identifiers only.", "text"),
        ("03-financial", "Client financial summary", "Personal monthly salary is $4,800. Savings at 20% are $960; annual savings are $11,520.", "finance"),
        ("04-incident", "Security incident report", "The technical discussion and example.invalid domain are demonstration content.", "image"),
        ("05-referral", "Synthetic referral letter", "Fictional patient reports a recurring headache. This is not a real medical record.", "medical"),
        ("06-expenses", "Expense reimbursement", "Review the attached synthetic screenshot and its cropped variant.", "crop"),
        ("07-research", "Research memo", "Public mathematical notation must remain intact during privacy review.", "math"),
        ("08-hidden-data", "Hidden-data stress document", "The current visible revision is public demonstration content.", "hidden"),
    ]
    records = []
    for identifier, title, description, kind in specs:
        path = output / (identifier + ".docx")
        if path.exists():
            raise FileExistsError(f"Refusing to overwrite {path}")
        doc = Document()
        doc.core_properties.author = "SafeShare synthetic fixtures"
        doc.core_properties.title = title
        doc.add_heading(title, 0)
        doc.add_paragraph(MARKER, "Subtitle")
        doc.add_paragraph(description)
        doc.sections[0].header.paragraphs[0].text = f"Test contact: {EMAIL}"
        doc.sections[0].footer.paragraphs[0].text = "Fictional content — not a real person or account"
        paragraph = doc.add_paragraph("Prepared for ")
        paragraph.add_run("Avery ").bold = True
        paragraph.add_run("Example").italic = True
        paragraph.add_run(". Identical names across runs must remain linked.")
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Field"
        table.rows[0].cells[1].text = "Value"
        for label, value in [("Name", NAME), ("Email", EMAIL), ("Reference", "DEMO-ONLY-0042")]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
        merged = table.add_row().cells
        merged[0].merge(merged[1]).text = "Merged table note: preserve this layout."
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("24486B")
        expected = [{"value": NAME, "category": "person", "action": "redact", "locations": ["body", "table"]},
                    {"value": EMAIL, "category": "email", "action": "redact", "locations": ["header", "table"]}]
        if kind == "finance":
            expected += [{"value": value, "category": "financial-derived" if value != "$4,800" else "financial", "action": "redact"}
                         for value in ["$4,800", "$960", "$11,520"]]
        if kind in ("image", "crop", "hidden"):
            doc.add_picture(str(image_path), width=Inches(6))
            if kind == "crop":
                # Original full pixels remain in the source DOCX beyond visible crop.
                pic = doc.inline_shapes[-1]._inline.graphic.graphicData.pic
                crop = OxmlElement("a:srcRect")
                crop.set("b", "50000")
                pic.blipFill.insert(1, crop)
            expected.append({"value": "DEMO_ONLY_NOT_A_REAL_SECRET", "category": "credential", "action": "redact", "locations": ["image"], "regions": regions})
        if kind == "math":
            equation(doc)
            doc.add_paragraph("Keep public equation (a+b)/c and mathematical structure.")
        if kind == "hidden":
            hidden_revision(doc)
            doc.core_properties.subject = "PRIVATE_METADATA_SENTINEL"
            expected += [{"value": "OLD_DELETED_PRIVATE_SENTINEL", "category": "revision", "action": "remove"},
                         {"value": "PRIVATE_METADATA_SENTINEL", "category": "metadata", "action": "remove"}]
        sentinel = "END_OF_DOCUMENT_" + identifier.upper().replace("-", "_")
        doc.add_paragraph(sentinel)
        doc.save(path)
        records.append({"id": identifier, "file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                        "language": "en", "source": "SafeShare-authored deterministic synthetic fixture; no external dataset records used",
                        "expected": expected, "preserve": [MARKER, sentinel, "Merged table note: preserve this layout."],
                        "features": {"tables": 1, "equations": int(kind == "math"), "images": int(kind in ("image", "crop", "hidden"))}})
    manifest = {"version": 1, "purpose": "Functional and demo fixtures, NOT an independent accuracy benchmark",
                "external_sources_used": [], "records": records}
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


def verify(output, manifest):
    for record in manifest["records"]:
        path = output / record["file"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == record["sha256"]
        with zipfile.ZipFile(path) as archive:
            body = ET.fromstring(archive.read("word/document.xml"))
            visible = "".join(body.itertext())
            assert record["preserve"][-2] in visible
            assert len(body.findall(".//w:tbl", NS)) == record["features"]["tables"]
            assert len(body.findall(".//m:oMath", NS)) == record["features"]["equations"]
            parts = "".join(ET.tostring(ET.fromstring(archive.read(name)), encoding="unicode")
                            for name in archive.namelist() if name.endswith(".xml"))
            assert EMAIL in parts
            assert NAME in visible
            for item in record["expected"]:
                if "image" not in item.get("locations", []):
                    assert item["value"] in parts, (record["id"], item["value"])
    assert len(manifest["records"]) == 8


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--self-check", action="store_true")
    args = parser.parse_args()
    if args.self_check:
        with tempfile.TemporaryDirectory(prefix="safeshare-word-") as directory:
            target = Path(directory)
            verify(target, build(target))
        print("8 synthetic English Word fixture checks passed.")
    elif args.output:
        verify(args.output, build(args.output))
        print(f"Created and verified 8 English DOCX fixtures: {args.output}")
    else:
        parser.error("Choose --output or --self-check")
