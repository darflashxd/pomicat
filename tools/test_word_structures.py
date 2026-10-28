"""Fictional DOCX structural tests in Chromium and Firefox. No network uploads."""
import base64
import hashlib
import io
import json
from pathlib import Path
import time
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from playwright.sync_api import sync_playwright
from audit_word_release import OUT, ROOT, SECRET

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def element(tag, text=None):
    node = OxmlElement(tag)
    if text is not None:
        node.text = text
    return node


def fixtures():
    cases = {}
    for kind in ['split-runs', 'merged-nested-table', 'content-control', 'equation-secret',
                 'equation-public', 'tracked-deletion', 'hidden-text', 'field-secret']:
        doc = Document()
        doc.add_paragraph('Fictional structural review. PUBLIC-SENTINEL-2026')
        p = doc.add_paragraph()
        if kind == 'split-runs':
            p.add_run(SECRET[:12]).bold = True
            p.add_run(SECRET[12:20]).italic = True
            p.add_run(SECRET[20:]).underline = True
        elif kind == 'merged-nested-table':
            table = doc.add_table(rows=3, cols=3)
            table.cell(0, 0).merge(table.cell(0, 2)).text = 'Merged heading'
            table.cell(1, 0).text = SECRET
            table.cell(1, 1).merge(table.cell(2, 1)).text = 'Vertical merge'
            nested = table.cell(2, 2).add_table(rows=1, cols=2)
            nested.cell(0, 0).text = SECRET
            nested.cell(0, 1).text = 'Public nested cell'
        elif kind == 'content-control':
            wrapper = element('w:sdt')
            content = element('w:sdtContent')
            paragraph = element('w:p')
            run = element('w:r'); run.append(element('w:t', SECRET))
            paragraph.append(run); content.append(paragraph); wrapper.append(content)
            p._p.addnext(wrapper)
        elif kind.startswith('equation'):
            p.add_run('Public equation: ' if kind == 'equation-public' else 'Private equation: ')
            math = element('m:oMath'); fraction = element('m:f')
            for tag, text in [('m:num', SECRET if kind == 'equation-secret' else '57600'), ('m:den', '12')]:
                term = element(tag); run = element('m:r'); run.append(element('m:t', text)); term.append(run); fraction.append(term)
            math.append(fraction); p._p.append(math)
        elif kind == 'tracked-deletion':
            deleted = element('w:del'); run = element('w:r'); run.append(element('w:delText', SECRET)); deleted.append(run); p._p.append(deleted)
        elif kind == 'hidden-text':
            run = p.add_run(SECRET); run.font.hidden = True
        elif kind == 'field-secret':
            field = element('w:fldSimple'); field.set(f'{{{W}}}instr', ' DOCVARIABLE ' + SECRET)
            run = element('w:r'); run.append(element('w:t', 'Public field display')); field.append(run); p._p.append(field)
        doc.sections[0].header.paragraphs[0].text = 'Public header'
        doc.sections[0].footer.paragraphs[0].text = 'Public footer'
        for run in doc.element.iter(f'{{{W}}}r'):
            properties = run.find(f'{{{W}}}rPr')
            if properties is None:
                properties = element('w:rPr'); run.insert(0, properties)
            color = element('w:color'); color.set(f'{{{W}}}val', '000000'); properties.append(color)
        data = io.BytesIO(); doc.save(data)
        cases[kind] = data.getvalue()
    return cases


def inspect(data):
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        xml = ET.fromstring(package.read('word/document.xml'))
        semantic = '\n'.join(''.join(ET.fromstring(package.read(name)).itertext())
                             for name in package.namelist() if name.endswith('.xml'))
        raw = b'\n'.join(package.read(name) for name in package.namelist() if name.endswith('.xml'))
        return {'secretInText': SECRET in semantic, 'secretInRaw': SECRET.encode() in raw,
                'tables': len(xml.findall('.//{' + W + '}tbl')),
                'merges': len(xml.findall('.//{' + W + '}gridSpan')) + len(xml.findall('.//{' + W + '}vMerge')),
                'equations': len(xml.findall('.//{' + M + '}oMath')),
                'controls': len(xml.findall('.//{' + W + '}sdt')),
                'publicPreserved': 'PUBLIC-SENTINEL-2026' in semantic}


def main():
    cases = fixtures()
    results = []
    with sync_playwright() as p:
        for engine in [p.chromium, p.firefox]:
            browser = engine.launch()
            page = browser.new_page(accept_downloads=True)
            page.goto('about:blank')
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            for script in ['redact.js', 'docx.js']:
                page.add_script_tag(content=(ROOT / 'extension/word' / script).read_text(encoding='utf-8'))
            for name, data in cases.items():
                (OUT / (name + '.docx')).write_bytes(data)
                start = time.monotonic()
                result = page.evaluate('''async ({data,secret}) => {
                  const bytes=Uint8Array.from(atob(data),c=>c.charCodeAt(0));
                  const doc=await inspectWordFile(bytes);
                  const visible=JSON.stringify(doc.blocks).includes(secret);
                  try {
                    const rules=[{value:secret,replacement:'[REDACTED]'}];
                    const output=await SafeShareWordRedact.sanitizeDocx(bytes.buffer,{rules});
                    const check=await SafeShareWordRedact.verifyDocx(output.bytes,rules);
                    let binary='';for(let i=0;i<output.bytes.length;i+=8192)binary+=String.fromCharCode(...output.bytes.subarray(i,i+8192));
                    window.auditBytes=output.bytes;
                    return {previewContainsSecret:visible,clean:check.clean,output:btoa(binary)};
                  } catch(error) {return {previewContainsSecret:visible,error:error.message};}
                }''', {'data': base64.b64encode(data).decode(), 'secret': SECRET})
                if 'output' in result:
                    output = base64.b64decode(result.pop('output'))
                    independent = inspect(output)
                    result['independent'] = independent
                    result['reopens'] = bool(Document(io.BytesIO(output)))
                    result['structurePreserved'] = all(independent[k] == inspect(data)[k] for k in ['tables','merges','equations','controls'])
                    result['falseClean'] = result['clean'] and (independent['secretInRaw'] or independent['secretInText'])
                    with page.expect_download() as event:
                        page.evaluate('''() => {const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([auditBytes]));a.download='reviewed.docx';a.click();}''')
                    downloaded = Path(event.value.path()).read_bytes()
                    result['downloadHashMatches'] = hashlib.sha256(downloaded).digest() == hashlib.sha256(output).digest()
                    (OUT / (engine.name + '-' + name + '-sanitized.docx')).write_bytes(output)
                result.update(browser=engine.name, case=name, seconds=round(time.monotonic()-start, 2))
                results.append(result)
            assert not errors, errors
            browser.close()
    (OUT / 'structural-results.json').write_text(json.dumps(results, indent=2), encoding='utf-8')
    print(json.dumps(results, indent=2))
    print('Report:', OUT / 'structural-results.json')
    # Report unsupported paths honestly instead of turning them into passing assertions.
    for result in results:
        if result['case'] in ['split-runs','merged-nested-table','content-control','equation-secret','equation-public','tracked-deletion','hidden-text']:
            assert result.get('clean') and not result.get('falseClean'), result
            assert result['structurePreserved'] and result['downloadHashMatches'], result
        if result['case'] == 'field-secret':
            assert 'field instruction' in result.get('error',''), result


if __name__ == '__main__':
    main()
