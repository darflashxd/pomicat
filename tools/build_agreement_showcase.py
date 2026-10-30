"""Build a fictional Word showcase. Run: python tools/build_agreement_showcase.py output.docx"""
import io
import sys
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAMES = ['Daniel James Carter', 'Melissa Jane Foster', 'Thomas Edward Brooks']

def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('233D50')
    section.header.paragraphs[0].text = 'NORTHBRIDGE SYSTEMS  /  PEOPLE & LEGAL'
    section.header.paragraphs[0].runs[0].font.size = Pt(9)
    section.footer.paragraphs[0].text = 'FICTIONAL SAMPLE — NOT FOR EXECUTION  |  Internal review copy'
    section.footer.paragraphs[0].runs[0].font.size = Pt(8)
    doc.add_paragraph('CONFIDENTIAL · DRAFT 03', 'Subtitle')
    doc.add_heading('Employee Separation\nAgreement', 0)
    doc.add_paragraph('Draft for Legal Review', 'Subtitle')
    doc.add_paragraph('Document NB-HR-2026-1042   •   Prepared 30 October 2026')
    doc.add_heading('01  Parties and effective date', 1)
    doc.add_paragraph('This draft records the proposed separation terms between Northbridge Systems '
        '(a fictional company) and the employee identified below. It is circulated to People Operations, '
        'Finance and Legal for review. No payment is authorized by this draft.')
    def table(rows):
        t = doc.add_table(rows=0, cols=2)
        t.style = 'Light Shading Accent 1'
        for left, right in rows:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = left, right
        return t
    table([('Employee', NAMES[0]), ('Position / employee ID', 'Senior Product Engineer / NB-2047'),
        ('Email', 'daniel.carter@example.invalid'), ('Home address', '1847 Willow Creek Avenue, Springfield, Illinois 62704'),
        ('Effective separation date', '30 November 2026')])
    doc.add_heading('02  Proposed settlement', 1)
    doc.add_paragraph('The monthly base salary is $7,850. The proposed separation payment is three '
        'months of base salary. Finance must confirm payroll treatment and any required deductions.')
    t = table([('Payment component', 'Gross amount (USD)'), ('Separation payment — 3 months', '$23,550'),
        ('Discretionary transition bonus', '$2,400'), ('Accrued leave payment', '$1,570'),
        ('Total proposed payment', '$27,520')])
    for row in list(t.rows)[1:]:
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph('Payment figures are subject to approval. The total is not a statement of net proceeds.')
    doc.add_page_break()
    doc.add_heading('03  Payment instructions', 1)
    doc.add_paragraph('Finance should verify the following fictional instructions through a separate '
        'approved channel before releasing any payment. Do not use this sample to send funds.')
    table([('Account holder', NAMES[0]), ('Bank', 'Example Payroll Bank — fictional'),
        ('Bank account number', '123456789012'), ('Target payment date', '15 December 2026')])
    doc.add_heading('04  Handover and company property', 1)
    doc.add_paragraph('The employee will return the assigned laptop and access card, complete the '
        'project handover, and confirm that company records remain in approved storage. Access '
        'removal will be coordinated by IT on the effective date. No passwords belong in this agreement.')
    doc.add_heading('05  Internal review notes', 1)
    doc.add_paragraph(f'{NAMES[1]}, People Operations, will confirm the employee details and the '
        'handover schedule. Legal must review the final wording before execution. Finance must '
        'reconcile the settlement table against payroll and obtain approval for the discretionary bonus.')
    doc.add_heading('06  Approval record', 1)
    table([('Reviewer', 'Status'), (f'People Operations — {NAMES[1]}', 'Pending'),
        (f'Finance — {NAMES[2]}', 'Pending'), ('Legal counsel', 'Pending legal review')])
    doc.add_paragraph('This sample is not legal advice and creates no agreement or obligation. All '
        'names, identifiers, account details and amounts are invented for a product demonstration.')
    doc.core_properties.author = 'Fictional showcase'
    doc.core_properties.title = 'Employee Separation Agreement — Fictional Sample'
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        text = archive.read('word/document.xml').decode()
        assert all(value in text for value in ['$23,550', '$2,400', '$1,570', '$27,520', NAMES[0]])
        assert 23550 + 2400 + 1570 == 27520
    return data

if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit('Usage: python tools/build_agreement_showcase.py output.docx')
    target = Path(sys.argv[1])
    if target.exists():
        raise SystemExit('Output already exists; choose a new path.')
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build())
    print(target)
