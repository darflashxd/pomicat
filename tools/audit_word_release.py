"""Local, fictional Word release audit. No relay uploads. Run with Python."""
import base64
import io
import json
import os
from pathlib import Path
import time
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import RGBColor
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get('LOCALAPPDATA', '.')) / 'Temp/opencode/word-release-audit'
OUT.mkdir(parents=True, exist_ok=True)
SECRET = 'SYNTHETIC-CANARY-ALPHA-73921'


def fixture():
    doc = Document()
    doc.add_heading('Restricted — Cross-border Acquisition and Incident Review', 0)
    doc.add_paragraph('Fictional audit material. No real people, accounts or credentials.')
    for chapter in range(45):
        doc.add_heading(f'{chapter + 1}. Regional review and approval record', 1)
        for paragraph in range(12):
            doc.add_paragraph(
                f'Record {chapter + 1}.{paragraph + 1}: The fictional Meridian board reviewed '
                'the acquisition schedule, regional tax treatment, employee transition and '
                'incident response obligations. Approval requires finance and legal review. '
                f'Private reference: {SECRET}. Public control: PUBLIC-SENTINEL-2026. '
                'Monthly salary USD 4,800; annual salary USD 57,600; bonus 20 percent. '
                'Contact: counsel@example.invalid. All amounts and identities are fictional.')
        table = doc.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells, ['Region', 'Monthly USD', 'Annual USD']):
            cell.text = text
        for region in ['Europe', 'Asia', 'Americas']:
            cells = table.add_row().cells
            for cell, text in zip(cells, [region, '4,800', '57,600']):
                cell.text = text
        doc.add_page_break()
    doc.sections[0].header.paragraphs[0].text = 'Restricted ' + SECRET
    doc.sections[0].footer.paragraphs[0].text = 'Review copy ' + SECRET
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
    stream = io.BytesIO()
    doc.save(stream)
    # python-docx's default template contains an unused bibliography customXml part.
    # Remove it with its references so the baseline exercises supported content.
    original = stream.getvalue()
    (OUT / 'default-template.docx').write_bytes(original)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as dest:
        for name in source.namelist():
            if name.startswith('customXml/'):
                continue
            content = source.read(name)
            if name.endswith('.rels') or name == '[Content_Types].xml':
                xml = ET.fromstring(content)
                for child in list(xml):
                    if 'customXml' in child.attrib.get('Type', '') or '/customXml/' in child.attrib.get('PartName', ''):
                        xml.remove(child)
                content = ET.tostring(xml, encoding='utf-8', xml_declaration=True)
            dest.writestr(name, content)
    return output.getvalue()


def repack(data, additions):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as dest:
        for name in source.namelist():
            dest.writestr(name, source.read(name))
        for name, content in additions.items():
            dest.writestr(name, content)
    return output.getvalue()


def main():
    data = fixture()
    # Realistic image-heavy size: valid, incompressible RGB PNGs, each below the part limit.
    from PIL import Image
    images = {}
    for i in range(4):
        png = io.BytesIO()
        Image.frombytes('RGB', (1000, 900), os.urandom(1000 * 900 * 3)).save(png, format='PNG')
        images[f'word/media/audit-{i}.png'] = png.getvalue()
    large_doc = Document(io.BytesIO(data))
    for content in images.values():
        large_doc.add_picture(io.BytesIO(content))
    large_stream = io.BytesIO()
    large_doc.save(large_stream)
    cases = {
        'long': data,
        'empty-bibliography': (OUT / 'default-template.docx').read_bytes(),
        '10mib-media': large_stream.getvalue(),
        'custom-properties': repack(data, {'docProps/custom.xml': f'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="2" name="PrivateReference"><vt:lpwstr>{SECRET}</vt:lpwstr></property></Properties>'}),
        'custom-xml': repack(data, {'customXml/item99.xml': f'<private>{SECRET}</private>'}),
        'oversized-inflation': repack(data, {'word/media/oversized.bin': b'A' * (9 * 1024 * 1024)}),
        'embedded-object': repack(data, {'word/embeddings/private.bin': SECRET.encode()}),
    }
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        parts = {n: source.read(n) for n in source.namelist()}
    parts['docProps/app.xml'] = parts['docProps/app.xml'].replace(b'<Company/>', ('<Company>' + SECRET + '</Company>').encode())
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    cases['company-metadata'] = stream.getvalue()
    results = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto('about:blank')
        for script in ['docx.js', 'redact.js']:
            page.add_script_tag(content=(ROOT / 'extension/word' / script).read_text(encoding='utf-8'))
        for name, content in cases.items():
            (OUT / (name + '.docx')).write_bytes(content)
            start = time.monotonic()
            result = page.evaluate('''async ({data, secret}) => {
              const bytes = Uint8Array.from(atob(data), c => c.charCodeAt(0));
              try {
                const doc = await inspectWordFile(bytes);
                const rules = [{value:secret, replacement:'[REDACTED]'}];
                const output = await SafeShareWordRedact.sanitizeDocx(bytes.buffer, {rules});
                const verification = await SafeShareWordRedact.verifyDocx(output.bytes, rules);
                let raw = '';
                for (let i=0;i<output.bytes.length;i+=8192) raw += String.fromCharCode(...output.bytes.subarray(i,i+8192));
                return {accepted:true,tables:doc.tables,images:doc.images,verificationClean:verification.clean,output:btoa(raw)};
              } catch(e) { return {accepted:false,error:e.message}; }
            }''', {'data': base64.b64encode(content).decode(), 'secret': SECRET})
            if 'output' in result:
                output = base64.b64decode(result.pop('output'))
                (OUT / (name + '-sanitized.docx')).write_bytes(output)
                with zipfile.ZipFile(io.BytesIO(output)) as package:
                    result['independentLeftovers'] = [n for n in package.namelist() if SECRET.encode() in package.read(n)]
                    result['publicSentinelPreserved'] = b'PUBLIC-SENTINEL-2026' in package.read('word/document.xml')
                result['outputBytes'] = len(output)
                reopened = Document(io.BytesIO(output))
                result['reopenedTables'] = len(reopened.tables)
                result['placedImages'] = len(reopened.inline_shapes)
            result.update(case=name, inputBytes=len(content), seconds=round(time.monotonic()-start, 2))
            results.append(result)
        browser.close()
    assert results[0]['verificationClean'] and not results[0]['independentLeftovers']
    assert results[0]['tables'] == 45 and results[0]['publicSentinelPreserved']
    bibliography = next(r for r in results if r['case'] == 'empty-bibliography')
    assert bibliography['accepted'] and bibliography['verificationClean'] and not bibliography['independentLeftovers']
    assert len(cases['10mib-media']) > 10 * 1024 * 1024
    assert next(r for r in results if r['case'] == '10mib-media')['placedImages'] == 4
    assert next(r for r in results if r['case'] == 'oversized-inflation')['accepted'] is False
    assert next(r for r in results if r['case'] == 'custom-xml')['accepted'] is False
    assert next(r for r in results if r['case'] == 'embedded-object')['accepted'] is False
    for name in ['custom-properties', 'company-metadata']:
        result = next(r for r in results if r['case'] == name)
        assert result['verificationClean'] and not result['independentLeftovers'], result
    (OUT / 'results.json').write_text(json.dumps(results, indent=2), encoding='utf-8')
    print(json.dumps(results, indent=2))
    print('Artifacts:', OUT)


if __name__ == '__main__':
    main()
