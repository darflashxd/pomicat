"""Build the fictional HR severance fixture used for Word parity checks."""
import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
TEXT = """Hey, can you help me sort this out before I talk to HR? My name is Daniel James Carter, and my email is daniel.carter@example.invalid. You can reach me at +1 202-555-0147.

My wife Emily Rose Carter and I live at 1847 Willow Creek Avenue, Springfield, Illinois 62704. My mom Patricia Anne Wilson lives in Portland, and my brother Michael David Carter is in Seattle.

My salary is $7,850 a month, my savings are $23,400, and my credit card debt is $6,750. My mortgage payment is $2,150 a month. My bank account number is 123456789012.

My manager Thomas Edward Brooks put me on a performance improvement plan last week. Melissa Jane Foster from HR mentioned a severance payment of $18,500.

My doctor Rebecca Louise Turner diagnosed me with type 2 diabetes. My latest medical bill is $1,840.

Also, I accidentally pasted my password, DemoOnly!NotReal4826, into our team notes. My database connection was in there too: postgresql://demo_user:FakePassword7319@db.internal.example:5432/payroll
"""


def black(run, size=None, bold=None):
    run.font.color.rgb = BLACK
    run.font.name = "Times New Roman"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    return run


def build(output: Path):
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise FileExistsError(f"Refusing to write into non-empty directory: {output}")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        black(footer.add_run("Internal draft - fictional data"), 8)
    black(doc.add_paragraph().add_run("Message to HR - working notes"), 16, True)
    for block in TEXT.split("\n\n"):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(10)
        black(paragraph.add_run(block), 11)
    black(doc.add_paragraph().add_run("Prepared with fictional details for testing."), 10)
    path = output / "HR severance letter.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if not args.output:
        parser.error("--output is required")
    print(build(args.output))
