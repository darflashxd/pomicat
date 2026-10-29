"""Build a fictional four-page confidential DOCX for regression tests.

All names, identifiers, accounts, amounts and credentials are invented. Never replace
this fixture with a real document: the point is a realistic shape, not real data.

Usage:
    python tools/build_confidential_fixture.py [output.docx]
"""
import io
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

SECRETS = {
    "names": ["Rina Kartika Wijaya", "Budi Santoso", "Sari Melati"],
    "nik": "3174092508910007",
    "npwp": "09.254.294.3-407.000",
    "address": "Jl. Cempaka Putih Tengah No. 42, Jakarta Pusat 10510",
    "phone": "+62 811-0000-0001",
    "email": "rina.wijaya@example.co.id",
    "bank": "5271 8890 1234",
    "amounts": [
        "Rp 28.500.000", "Rp 4.275.000", "Rp 7.125.000", "Rp 42.750.000",
        "Rp 9.800.000", "Rp 2.350.000", "Rp 342.000.000", "Rp 118.000.000",
    ],
    "health": ["anxiety disorder", "Lexapro 20 mg"],
    "password": "Nusantara#2026!",
    "dsn": "postgres://hr_admin:Nusantara%232026@db.internal.example:5432/hr",
}


def _logo_bytes():
    image = Image.new("RGB", (240, 72), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle([0, 0, 239, 71], outline=(20, 40, 80), width=3)
    draw.text((18, 18), "NUSANTARA", fill=(20, 40, 80))
    draw.text((18, 40), "TEKNOLOGI", fill=(120, 30, 30))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _page_field(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    paragraph._p.append(field)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    header = section.header.paragraphs[0]
    header.text = "STRICTLY CONFIDENTIAL — INTERNAL HR USE ONLY"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(0x99, 0x00, 0x00)
    header.add_run().add_picture(io.BytesIO(_logo_bytes()), width=Inches(1.1))

    footer = section.footer.paragraphs[0]
    footer.text = "PT Nusantara Teknologi — Confidential Severance Review — Page "
    _page_field(footer)
    footer.runs[0].font.size = Pt(8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CONFIDENTIAL MEMORANDUM")
    run.bold = True
    run.font.size = Pt(16)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Employee Severance and Separation Review").italic = True

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    for index, (key, value) in enumerate([
        ("To", "Board of Directors, People Committee"),
        ("From", "Sari Melati — Head of People Operations"),
        ("Date", "29 October 2026"),
        ("Reference", "NT-HR-2026-1042 / CONFIDENTIAL"),
    ]):
        meta.cell(index, 0).text = key
        meta.cell(index, 1).text = value

    doc.add_paragraph("1. Employee identity")
    ident = doc.add_table(rows=7, cols=2)
    ident.style = "Table Grid"
    for index, (key, value) in enumerate([
        ("Full name", SECRETS["names"][0]),
        ("Employee ID", "NT-2021-0457"),
        ("NIK", SECRETS["nik"]),
        ("NPWP", SECRETS["npwp"]),
        ("Home address", SECRETS["address"]),
        ("Mobile", SECRETS["phone"]),
        ("Email", SECRETS["email"]),
    ]):
        ident.cell(index, 0).text = key
        ident.cell(index, 1).text = value

    doc.add_paragraph(
        "The employee has requested a written explanation of the proposed separation package "
        "before signing any agreement. This memorandum is prepared for the People Committee and "
        "legal counsel only."
    )

    doc.add_page_break()
    doc.add_paragraph("2. Compensation and proposed package")
    comp = doc.add_table(rows=7, cols=3)
    comp.style = "Table Grid"
    rows = [
        ("Component", "Monthly (IDR)", "Notes"),
        ("Base salary", SECRETS["amounts"][0], "Current"),
        ("Fixed allowance", SECRETS["amounts"][1], "Transport and housing"),
        ("Annual bonus (2025)", SECRETS["amounts"][2], "Paid in January"),
        ("Proposed severance", SECRETS["amounts"][3], "1.5 months base"),
        ("Mortgage instalment", SECRETS["amounts"][4], "KPR, private"),
        ("Credit card minimum", SECRETS["amounts"][5], "Private"),
    ]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            comp.cell(row_index, column_index).text = value
    doc.add_paragraph(
        f"Annualised base salary is 12 x {SECRETS['amounts'][0]} = {SECRETS['amounts'][6]}. "
        f"After the proposed severance and outstanding mortgage, the projected remaining runway "
        f"is approximately {SECRETS['amounts'][7]} before tax."
    )
    doc.add_paragraph(
        f"Payroll account for the final settlement: BCA {SECRETS['bank']} "
        f"(a/n {SECRETS['names'][0]})."
    )

    doc.add_page_break()
    doc.add_paragraph("3. Health and performance context")
    doc.add_paragraph(
        "Medical note: the employee disclosed a diagnosis of generalized anxiety disorder and "
        "currently takes Lexapro 20 mg each night. Treat this information as sensitive health data; "
        "do not circulate outside HR and the company doctor."
    )
    doc.add_paragraph(
        "Performance note: the manager, Budi Santoso, placed the employee on a performance "
        "improvement plan (PIP) in July 2026. The PIP review was delayed and the employee disputes "
        "the timeline."
    )
    doc.add_paragraph(
        f"Internal staging credentials referenced during the audit: staging password "
        f"{SECRETS['password']} and {SECRETS['dsn']} — rotate before any further testing."
    )
    hidden = doc.add_paragraph(
        "Hidden note: do not include the employee's therapy details in the board pack."
    )
    hidden.runs[0].font.hidden = True

    doc.add_page_break()
    doc.add_paragraph("4. Legal terms and acknowledgement")
    doc.add_paragraph(
        "The employee agrees that the separation is voluntary, that the final settlement is paid "
        "within 14 working days, and that the confidentiality obligation survives termination. "
        "Nothing in this memorandum prevents the employee from reporting unlawful conduct to a "
        "regulator."
    )
    doc.add_paragraph(
        "Reviewers: Sari Melati (People Operations), Budi Santoso (Engineering Manager), Legal Counsel."
    )
    signature = doc.add_table(rows=3, cols=2)
    signature.style = "Table Grid"
    for index, (key, value) in enumerate([
        ("Prepared by", "Sari Melati"),
        ("Reviewed by", "Legal Counsel"),
        ("Acknowledged by", "Rina Kartika Wijaya"),
    ]):
        signature.cell(index, 0).text = key
        signature.cell(index, 1).text = value
    doc.add_paragraph(
        "Distribution: Board People Committee, Legal Counsel, Head of People Operations. Do not forward."
    )

    properties = doc.core_properties
    properties.author = "HR Department"
    properties.title = "Confidential Severance Review"
    properties.subject = "NT-HR-2026-1042"
    properties.comments = "Internal only. Fictional fixture for testing."

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    data = build()
    target = sys.argv[1] if len(sys.argv) > 1 else "confidential-fixture.docx"
    with open(target, "wb") as handle:
        handle.write(data)
    print(f"{target}: {len(data)} bytes")
